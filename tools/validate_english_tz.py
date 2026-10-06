from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def collect_text(document: Document) -> str:
    chunks: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    for section in document.sections:
        chunks.extend(paragraph.text for paragraph in section.header.paragraphs)
        chunks.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(chunks)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()

    document = Document(args.docx)
    text = collect_text(document)

    required = [
        "Kids → A1 → A2 → B1 → B2",
        "финальный Kids; остановить тест после 14",
        "ONE_LEVEL_DOWNGRADE_FLOOR",
        "downgradeFloorApplied",
        "Завершить тест без результата?",
        "Уровень пока не определён",
        "STUDENT_EXITED_BEFORE_COMPLETION",
        "Назад",
        "1518×659 и 390×844",
        "The lift is out of order. Please use the stairs. What should people do?",
        "ТЗ v4.0",
    ]
    forbidden = [
        "После отправки ответа нельзя вернуться и изменить его.",
        "Каждая завершённая попытка содержит ровно 40 заданий",
        "пересчитать все 40 ответов и опуститься ниже",
        "Сценарий 1 — итог Pre-A1",
        "Всего 40 заданий в трёх частях.",
    ]

    missing = [value for value in required if value not in text]
    stale = [value for value in forbidden if value in text]
    if missing or stale:
        raise SystemExit(f"missing={missing}; stale={stale}")

    print(
        f"OK paragraphs={len(document.paragraphs)} tables={len(document.tables)} "
        f"sections={len(document.sections)}"
    )


if __name__ == "__main__":
    main()
