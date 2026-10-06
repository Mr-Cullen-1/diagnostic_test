from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def remove_element(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def set_paragraph_text(paragraph, text: str) -> None:
    """Replace text while retaining the first run's direct formatting."""
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in list(paragraph.runs[1:]):
            remove_element(run._r)
    else:
        paragraph.add_run(text)


def set_label_body(paragraph, label: str, body: str) -> None:
    """Retain the paragraph's existing label/body run formatting."""
    runs = paragraph.runs
    if len(runs) >= 2:
        label_run = runs[0]
        body_run = runs[1]
        label_run.text = label
        body_run.text = body
        for run in list(paragraph.runs[2:]):
            remove_element(run._r)
        return

    template_rpr = deepcopy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else None
    set_paragraph_text(paragraph, label)
    label_run = paragraph.runs[0]
    label_run.bold = True
    body_run = paragraph.add_run(body)
    if template_rpr is not None:
        body_run._r.insert(0, deepcopy(template_rpr))
    body_run.bold = False


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in list(cell.paragraphs[1:]):
        remove_element(extra._p)


def clone_paragraph_after(paragraph, text: str) -> Paragraph:
    cloned_xml = deepcopy(paragraph._p)
    paragraph._p.addnext(cloned_xml)
    cloned = Paragraph(cloned_xml, paragraph._parent)
    set_paragraph_text(cloned, text)
    return cloned


def clone_row_before(table, before_index: int, source_index: int):
    cloned_xml = deepcopy(table.rows[source_index]._tr)
    table.rows[before_index]._tr.addprevious(cloned_xml)
    return table.rows[before_index]


def clone_row_at_end(table, source_index: int):
    cloned_xml = deepcopy(table.rows[source_index]._tr)
    table._tbl.append(cloned_xml)
    return table.rows[-1]


def set_row(row, values: list[str]) -> None:
    for cell, value in zip(row.cells, values, strict=True):
        set_cell_text(cell, value)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def update_document(input_path: Path, output_path: Path) -> None:
    document = Document(input_path)
    paragraphs = document.paragraphs
    tables = document.tables

    # Cover and metadata.
    set_paragraph_text(
        paragraphs[2],
        "Точная логика определения уровня: Kids → A1 → A2 → B1 → B2",
    )
    set_paragraph_text(
        paragraphs[3],
        "= 14 вопросов для результата Kids или 40 для полного маршрута",
    )
    set_paragraph_text(tables[0].cell(0, 0).paragraphs[1], "старт для всех")
    set_paragraph_text(tables[0].cell(0, 2).paragraphs[1], "адаптация")
    set_paragraph_text(tables[0].cell(0, 4).paragraphs[1], "подтверждение")
    set_cell_text(tables[1].cell(0, 1), "4.0")
    set_cell_text(tables[1].cell(1, 1), "3 августа 2026")
    set_paragraph_text(
        tables[2].cell(0, 0).paragraphs[1],
        "Ученик не выбирает свой уровень. После 14 стартовых вопросов система либо завершает тест одним результатом Kids, либо проверяет границу знаний ещё в двух частях и рекомендует курс после 40-го ответа.",
    )

    # Product-facing labels and course routing.
    set_cell_text(
        tables[3].cell(2, 1),
        "Не влияет на placementLevel. Для полного маршрута с внутренним preA1 выбирает Kids (до 10 лет) или A1 Foundation (старше); ранний Kids после части 1 всегда рекомендует Kids.",
    )
    set_paragraph_text(
        paragraphs[12],
        "Kids — пользовательское название результата с внутренним ключом preA1; текст «Pre-A1» в интерфейсе не показывается.",
    )
    set_cell_text(tables[5].cell(1, 0), "Kids (внутри: preA1)")
    set_cell_text(
        tables[5].cell(1, 2),
        "После части 1 → Kids для любого возраста. После полного маршрута: до 10 лет → Kids; старше → A1 Foundation.",
    )

    # Conditional three-part flow.
    set_paragraph_text(
        tables[6].cell(0, 1).paragraphs[1],
        "7 заданий A1 + 7 заданий A2. Если A1 = 0–2/7, тест заканчивается одним итогом Kids. Иначе сразу начинается часть 2.",
    )
    set_paragraph_text(
        tables[7].cell(0, 1).paragraphs[1],
        "4 мини-блока × 4. Запускается только после результата A1 или выше на старте. После вопроса 30 система сразу открывает часть 3 без экрана результата блока.",
    )
    set_paragraph_text(
        tables[8].cell(0, 1).paragraphs[1],
        "Система выбирает две соседние внутренние ступени: preA1/A1, A1/A2, A2/B1 или B1/B2. В интерфейсе preA1 называется Kids.",
    )
    set_paragraph_text(
        tables[9].cell(0, 1).paragraphs[1],
        "5 вопросов нижнего уровня + 5 вопросов верхнего уровня. После вопроса 40 итоговый экран появляется сразу, без отдельного результата части.",
    )
    set_paragraph_text(
        tables[10].cell(0, 1).paragraphs[1],
        "Kids, A1, A2, B1 или B2 + уверенность high/medium + возможная пограничная отметка. Итог показывается один раз.",
    )
    set_paragraph_text(
        tables[11].cell(0, 0).paragraphs[1],
        "Старт ищет район и имеет одно терминальное исключение: внутренний preA1 сразу даёт Kids после 14 вопросов. Во всех остальных случаях адаптация двигает тест, а подтверждение принимает решение. Промежуточные результаты частей не показываются.",
    )
    set_paragraph_text(
        paragraphs[23],
        "Все ученики начинают одинаково. Обычно эта часть выбирает, с какой сложности начать адаптацию. Исключение: если A1 набрано 0–2/7, тест завершается после 14 вопросов одним валидным результатом Kids.",
    )
    set_cell_text(tables[13].cell(0, 2), "Маршрут после старта")
    set_cell_text(
        tables[13].cell(1, 2),
        "Внутренний preA1 → финальный Kids; остановить тест после 14",
    )
    set_paragraph_text(
        tables[14].cell(0, 0).paragraphs[0],
        "Два критических правила после старта",
    )
    set_paragraph_text(
        tables[14].cell(0, 0).paragraphs[1],
        "Если A1 = 0–2/7, система сразу показывает один итог Kids и не запускает части 2–3. Даже 14/14 не даёт B2 автоматически: старт содержит только A1 и A2, поэтому сильный ученик начинает адаптацию с B1 и доказывает B2 дальше.",
    )

    # Direct transitions and final confirmation rules.
    set_paragraph_text(
        paragraphs[29],
        "Один мини-блок — это четыре задания, после которых система один раз решает: сделать следующий блок легче, оставить таким же или сделать сложнее. По завершении всех 16 заданий часть 3 открывается сразу; отдельного экрана «Результат блока» нет.",
    )
    set_paragraph_text(
        paragraphs[45],
        "Последняя часть проверяет ровно ту границу, возле которой оказался ученик. После десятого ответа система сразу показывает итоговый Placement; дополнительного экрана результата части и кнопки «Показать итог» нет.",
    )
    set_cell_text(
        tables[25].cell(4, 2),
        "Выбранная пара была слишком высокой: итог = нижний уровень пары; confidence = medium; inconsistency = true",
    )
    set_paragraph_text(
        tables[26].cell(0, 0).paragraphs[0],
        "Почему 3/5 и слабая пара не обнуляют маршрут",
    )
    set_paragraph_text(
        tables[26].cell(0, 0).paragraphs[1],
        "Три правильных ответа из пяти — это граница, поэтому ставится нижний уровень и сохраняется borderlineWith. Если обе стороны подтверждения слабы, система также выбирает нижний уровень пары с medium/inconsistency, а не пересчитывает результат на несколько ступеней вниз. После этого действует общий floor: финальный уровень не может быть ниже чем на одну ступень от последнего probeLevel части 2 (B2 → минимум B1; B1 → A2; A2 → A1; A1 → внутренний preA1/Kids).",
    )
    prevent_row_split(tables[26].rows[0])

    # Replace obsolete Pre-A1 full-route scenario with the valid short Kids route.
    kids_cell = tables[27].cell(0, 0)
    kids_paragraphs = kids_cell.paragraphs
    set_paragraph_text(kids_paragraphs[0], "Сценарий 1 — итог Kids после части 1")
    set_label_body(
        kids_paragraphs[1],
        "Скрытая реальность: ",
        "Ученик узнаёт отдельные слова, но база A1 пока не держится.",
    )
    set_label_body(
        kids_paragraphs[2],
        "Старт 14: ",
        "A1 = 1/7, A2 = 0/7 → внутренний probeLevel = preA1.",
    )
    set_label_body(
        kids_paragraphs[8],
        "Дальнейшие части: ",
        "Адаптация и подтверждение не запускаются; нет промежуточного результата и нет второй кнопки показа итога.",
    )
    set_label_body(
        kids_paragraphs[9],
        "Результат: ",
        "Kids, recommendedCourse = Kids, answeredCount = totalQuestions = 14. Экран результата показывается один раз.",
    )
    for obsolete in list(kids_paragraphs[3:8]):
        remove_element(obsolete._p)

    # Add a concrete one-level-floor regression to the existing scenario block.
    floor_cell = tables[32].cell(0, 0)
    floor_template = floor_cell.paragraphs[-1]
    floor_note = clone_paragraph_after(
        floor_template,
        "Регрессия B2-floor: если preConfirmationProbe = B2, а сырой итог подтверждения ниже B1, финальный placement = B1; confidence = medium; downgradeFloorApplied = true.",
    )
    floor_note.runs[0].bold = True
    set_cell_text(tables[33].cell(1, 0), "preA1 / A1 (внутри)")
    set_cell_text(tables[33].cell(1, 2), "Kids, близко к A1")
    set_paragraph_text(
        tables[34].cell(0, 0).paragraphs[1],
        "В пограничном случае курс выбирается по нижнему уровню пары, а следующая ступень сохраняется как borderlineWith. При этом финальное понижение ограничено одной ступенью относительно последнего probeLevel адаптации: достигнутый в части 2 B2 не может превратиться в A2 или Kids; минимум — B1.",
    )

    # File structure, bank language, and exact corrected item.
    set_cell_text(tables[35].cell(1, 0), "english-diagnostic(1).jsx")
    set_cell_text(tables[35].cell(3, 0), "english-test-items-a2-b1(1).jsx")
    total_index = len(tables[35].rows) - 1
    policy_row = clone_row_before(tables[35], total_index, 4)
    set_row(
        policy_row,
        [
            "english-placement-policy.js",
            "0 вопросов",
            "Чистая политика ограничения понижения на одну ступень",
        ],
    )
    total_index = len(tables[35].rows) - 1
    tests_row = clone_row_before(tables[35], total_index, 4)
    set_row(
        tests_row,
        [
            "tests/english-placement-policy.test.js",
            "4 теста",
            "Регрессии floor, включая B2 → минимум B1",
        ],
    )
    set_cell_text(
        tables[35].rows[-1].cells[2],
        "Общий банк 150; ученик видит 14 (Kids) или 40 (полный маршрут)",
    )
    set_paragraph_text(
        paragraphs[68],
        "Итоговое распределение в полном маршруте из 40. Для валидного раннего Kids используется только старт из 14:",
    )
    set_paragraph_text(
        paragraphs[72],
        "В данных используется level: \"preA1\", а не level: \"kids\". Пользовательский интерфейс всегда преобразует preA1 в Kids.",
    )
    set_paragraph_text(
        paragraphs[75],
        "Задания с внутренним level: \"preA1\" должны быть простыми, но не выглядеть детскими: их может увидеть подросток или взрослый новичок.",
    )
    item_note = clone_paragraph_after(
        paragraphs[76],
        "Контрольный item A1-r-07: prompt = “The lift is out of order. Please use the stairs. What should people do?”; правильный ответ = “Use the stairs”.",
    )
    item_note.style = paragraphs[76].style

    # Interface behavior: progress, navigation, termination, and no-scroll design.
    set_paragraph_text(
        tables[39].cell(0, 0).paragraphs[1],
        "Узнай, с какого курса лучше начать. Тест сам подберёт сложность. До 40 заданий в трёх частях. [Начать тест]",
    )
    set_paragraph_text(
        paragraphs[84],
        "Показывать: «Блок 2 из 3» и общий номер: «Задание 19 из 40».",
    )
    set_paragraph_text(
        paragraphs[88],
        "Кнопка «Назад» разрешает вернуться к предыдущему вопросу только внутри текущей части, восстановить ответ и изменить его. Переход назад через границы 14→15 и 30→31 запрещён; при изменении ответа зависимая адаптивная ветка пересчитывается, а проверка звука, число прослушиваний и audioError не сбрасываются.",
    )
    set_paragraph_text(
        paragraphs[89],
        "Автоматическое прерывание сохраняет активный прогресс для продолжения. Осознанное «Завершить тест» после подтверждения создаёт терминальное состояние incomplete и предлагает пройти тест заново.",
    )
    set_paragraph_text(
        paragraphs[90],
        "Непоказанные вопросы не считаются ошибками. В incomplete нет placementLevel, результата, балла, статистики навыков или рекомендации курса.",
    )
    ui_anchor = paragraphs[90]
    ui_requirements = [
        "В части 1 сверху показывается ровно один прогресс-бар. После прохода выше Kids показываются только два бара — для частей 2 и 3. На финальном и incomplete-экранах баров нет.",
        "После частей 1 (если не Kids) и 2 переходить прямо к следующему вопросу. Не показывать правильные/неправильные ответы, уровень или экран «Результат блока».",
        "Исключение — Kids после части 1: это валидный финал, который сразу показывается одним экраном и не дублируется.",
        "Кнопка «Завершить тест» открывает подтверждение «Завершить тест без результата?» и предупреждение: «Если завершить тест сейчас, уровень не будет определён. Чтобы получить рекомендацию, нужно ответить на все вопросы». После подтверждения показать «Тест не завершён», «Уровень пока не определён», N из 40 и предложение пройти заново.",
        "Финальный результат использует компактную сводку уровня/курса, две компактные карточки навыков и всегда видимые действия «Пройти заново» и «К выбору курса».",
        "Вопрос, проверка звука, итоговый результат и incomplete должны помещаться без вертикальной прокрутки страницы и внутренних контейнеров. Нижние кнопки всегда видимы; контрольные размеры — 1518×659 и 390×844.",
        "Схема сохранённой сессии — version 6. Активная попытка и incomplete восстанавливаются после перезагрузки; retry создаёт чистую попытку.",
    ]
    for requirement in ui_requirements:
        ui_anchor = clone_paragraph_after(ui_anchor, requirement)

    set_paragraph_text(
        tables[40].cell(0, 0).paragraphs[1],
        "Нельзя копировать логику, где сложность меняется после каждого отдельного ответа или где при досрочном завершении все оставшиеся вопросы записываются как ошибки. В English Placement решение меняется только после мини-блока из 4. Ручное завершение даёт incomplete без уровня, а валидный Kids после 14 — один настоящий результат.",
    )

    # Updated pseudocode and API shapes.
    pseudocode = """start = getStartProbe(14 answers)
if start.probeLevel == preA1:
  return complete(level=Kids, course=Kids, total=14)

for block in 1..4:
  items = buildMiniBlock(probeLevel, 4)
  correct = score(items)
  if correct >= 3: probeLevel = oneLevelUp(probeLevel)
  if correct == 2: probeLevel = probeLevel
  if correct <= 1: probeLevel = oneLevelDown(probeLevel)

preConfirmationProbe = probeLevel
stats30 = calculateLevelStats(first30Responses)
supported = highestLevel(seen >= 4, accuracy >= 60%, lower >= 70%)
pair = confirmationPair(supported)
confirm = show(5 lower + 5 upper)
raw = decideFromConfirmation(confirm, all40)
floor = oneLevelDown(preConfirmationProbe)
result = clampNotBelow(raw, floor)

if studentFinishesEarly:
  return incomplete(result=null, answeredCount, total=40)"""
    set_paragraph_text(tables[41].cell(0, 0).paragraphs[0], pseudocode)

    payload = """complete = {
  status: "complete",
  completionType: "FULL_PLACEMENT", // or SECTION_1_KIDS
  placementLevel: "A2",             // preA1 displays as Kids
  recommendedCourse: "A2",
  confidence: "medium",
  borderlineWith: "B1",
  answeredCount: 40,                 // 14 for SECTION_1_KIDS
  preConfirmationProbe: "B1",
  calculatedPlacementLevel: "A2",
  downgradeFloor: "A2",
  downgradeFloorApplied: false,
  reasonCode: "BORDERLINE_UPPER_3_OF_5",
  // if floor applies: reasonCode = "ONE_LEVEL_DOWNGRADE_FLOOR"
  adaptivePath: ["A2:S", "B1:M", "B1:W", "A2:S"],
  startStats: { A1: "7/7", A2: "5/7" },
  confirmStats: { A2: "5/5", B1: "3/5" }
}

incomplete = {
  phase: "incomplete",
  status: "incomplete",
  result: null,
  reasonCode: "STUDENT_EXITED_BEFORE_COMPLETION",
  answeredCount: 17,
  totalQuestions: 40,
  endedAt: "ISO timestamp"
}"""
    set_paragraph_text(tables[42].cell(0, 0).paragraphs[0], payload)

    # Plain-language decisions.
    for label, explanation in [
        (
            "Один бар, затем два",
            "Сначала виден только путь первой части. После прохода выше Kids ученик видит две оставшиеся части, не три бара сразу.",
        ),
        (
            "Назад только внутри части",
            "Ответ можно пересмотреть, но уже завершённая часть остаётся закрытой, чтобы не ломать подтверждённый маршрут.",
        ),
        (
            "Без экранов результата частей",
            "Промежуточные баллы не отвлекают и не подсказывают уровень; исключение — настоящий финал Kids после старта.",
        ),
    ]:
        row = clone_row_at_end(tables[43], len(tables[43].rows) - 1)
        set_row(row, [label, explanation])

    # Risks and readiness.
    set_cell_text(tables[45].cell(1, 0), "14 или 40 вопросов ≠ автоматически точный тест")
    floor_risk = clone_row_at_end(tables[45], len(tables[45].rows) - 1)
    set_row(
        floor_risk,
        [
            "Ограничение понижения на одну ступень — продуктовый guardrail",
            "Случайный выход на высокий probe может удержать ученика выше, чем показывает слабое подтверждение",
            "Логировать downgradeFloorApplied и сверять с преподавателем и успехом на курсе",
        ],
    )

    set_paragraph_text(
        paragraphs[109],
        "□ Валидная завершённая попытка содержит 14 заданий для раннего Kids или 40 заданий для полного маршрута; ручное завершение — incomplete, а не результат.",
    )
    set_paragraph_text(
        paragraphs[121],
        "□ Непоказанные вопросы не считаются ошибками; ручное завершение не выдаёт уровень, балл или курс.",
    )
    set_paragraph_text(
        paragraphs[122],
        "□ Возраст не влияет на placementLevel. Ранний Kids после части 1 всегда рекомендует Kids; для полного маршрута с внутренним preA1 возраст выбирает Kids/A1 Foundation.",
    )
    set_paragraph_text(
        paragraphs[123],
        "□ Пользовательские результаты ограничены Kids, A1, A2, B1, B2; внутренний алгоритм и банк используют ключ preA1.",
    )
    set_paragraph_text(
        paragraphs[124],
        "□ Требуются unit/e2e-сценарии для Kids, A1, A2, B1, B2, четырёх пограничных пар и ограничения понижения.",
    )
    dod_anchor = paragraphs[124]
    dod_items = [
        "□ При probeLevel = B2 после части 2 финальный уровень не опускается ниже B1; аналогичный floor действует для остальных ступеней.",
        "□ После частей нет промежуточных экранов результата; ранний Kids и финал после вопроса 40 показываются сразу и по одному разу.",
        "□ В части 1 виден один прогресс-бар, затем — два бара для частей 2 и 3.",
        "□ «Назад» восстанавливает и изменяет ответ внутри части, но недоступен на первом вопросе следующей части.",
        "□ Подтверждённое ручное завершение показывает incomplete/unknown и не вызывает completion callback.",
        "□ Item A1-r-07 содержит точный текст про lift/stairs и правильный ответ «Use the stairs».",
        "□ На контрольных desktop/mobile viewport нет вертикального scroll, обрезанных вопросов или недоступных нижних кнопок.",
    ]
    for item in dod_items:
        dod_anchor = clone_paragraph_after(dod_anchor, item)

    set_paragraph_text(
        tables[46].cell(0, 0).paragraphs[1],
        "Разработчик может прогнать все сценарии и регрессии из этого ТЗ: короткий Kids, полный маршрут, incomplete, пограничные пары, back-навигацию, one-level floor и no-scroll layout. Методолог может открыть лог и понять выбор пары, сырой результат и применённый downgrade floor.",
    )
    set_paragraph_text(
        paragraphs[131],
        "Примечание: конкретные числа 14 + 16 + 10, пороги 3/4, 60%/70%, 4/5 и 65%, раннее завершение Kids и ограничение понижения на одну ступень являются продуктовыми гипотезами этого ТЗ. Их нужно подтвердить пилотом; это не готовый официальный стандарт Cambridge.",
    )

    # Footer version/count contract.
    for section in document.sections:
        for footer_paragraph in section.footer.paragraphs:
            if "English Placement Test" in footer_paragraph.text:
                for run in footer_paragraph.runs:
                    if "English Placement Test" in run.text:
                        run.text = "English Placement Test  •  ТЗ v4.0  •  14 или 40\t"
                        break

    # Longer v4 sections can flow onto the next page. Let the following
    # headings use the available continuation page instead of forcing sparse
    # pages with only one or two carry-over paragraphs.
    for paragraph in document.paragraphs:
        if paragraph.text.startswith(
            (
                "6. Понятные сценарии",
                "8. Требования к интерфейсу",
                "11. Критерии готовности",
            )
        ):
            paragraph.paragraph_format.page_break_before = False

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()
    update_document(args.input_docx, args.output_docx)


if __name__ == "__main__":
    main()
