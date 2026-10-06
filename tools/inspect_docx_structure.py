from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--start", type=int, default=0)
    parser.add_argument("--count", type=int, default=None)
    parser.add_argument("--tables-only", action="store_true")
    parser.add_argument("--skip-tables", action="store_true")
    parser.add_argument("--detail-table", type=int, action="append", default=[])
    parser.add_argument("--headers", action="store_true")
    args = parser.parse_args()

    document = Document(args.docx)
    if args.headers:
        for section_index, section in enumerate(document.sections):
            print(f"SECTION {section_index}")
            for name, container in (("HEADER", section.header), ("FOOTER", section.footer)):
                for paragraph_index, paragraph in enumerate(container.paragraphs):
                    print(
                        f"{name} P{paragraph_index} style={paragraph.style.name} "
                        f"text={paragraph.text!r}"
                    )
                for table_index, table in enumerate(container.tables):
                    print(f"{name} TABLE {table_index} text={table.cell(0, 0).text!r}")
    if not args.tables_only:
        stop = None if args.count is None else args.start + args.count
        print(f"PARAGRAPHS {len(document.paragraphs)} showing={args.start}:{stop}")
        for index, paragraph in enumerate(document.paragraphs[args.start:stop], args.start):
            text = paragraph.text.replace("\n", "\\n")
            print(f"P{index:03d}\t{paragraph.style.name}\t{text}")

    if not args.skip_tables:
        print(f"TABLES {len(document.tables)}")
        for table_index, table in enumerate(document.tables):
            print(f"TABLE {table_index} rows={len(table.rows)} cols={len(table.columns)}")
            for row_index, row in enumerate(table.rows):
                cells = [cell.text.replace("\n", "\\n") for cell in row.cells]
                print(f"T{table_index}R{row_index:03d}\t" + " || ".join(cells))
                if table_index in args.detail_table:
                    for cell_index, cell in enumerate(row.cells):
                        for paragraph_index, paragraph in enumerate(cell.paragraphs):
                            runs = [
                                {
                                    "text": run.text,
                                    "bold": run.bold,
                                    "italic": run.italic,
                                    "size": run.font.size.pt if run.font.size else None,
                                    "color": str(run.font.color.rgb) if run.font.color.rgb else None,
                                }
                                for run in paragraph.runs
                            ]
                            print(
                                f"  C{cell_index}P{paragraph_index} "
                                f"style={paragraph.style.name} runs={runs}"
                            )


if __name__ == "__main__":
    main()
